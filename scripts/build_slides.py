#!/usr/bin/env python3
"""Gera a apresentacao do WPMC 2026 (22 slides + 4 de reserva) em .pptx editavel.

Plano aprovado em 22/10 (storytelling "a pergunta que os revisores fizeram"). O
visual segue paper/slides/WirelessDays-Enoki-2025 (reference).pdf, a apresentacao
de referencia do grupo: fundo #F9F9F9, faixa cinza inferior, linhas finas,
titulo centralizado, marcadores quadrados, fontes Commissioner (capa) e Golos Text
(corpo). As duas fontes sao Google Fonts, entao o arquivo abre com elas no Google
Slides.

Os numeros dos slides sao CALCULADOS dos mesmos dados que o
scripts/verify_paper_numbers.py confere contra o artigo (commit e4de793), nao
digitados. Os graficos vem de scripts/build_slide_figures.py.

Uso:
    venv/bin/python scripts/build_slide_figures.py
    venv/bin/python scripts/build_slides.py
Saida: paper/slides/WPMC2026-presentation.pptx
Logos opcionais em paper/slides/logos/*.png|jpg (sem eles, fica um espaco reservado).
"""

from __future__ import annotations

import glob
import json
import math
import os
import re
import sys

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "scripts"))
from analyze_results import (compliance_table, latency_stats, load_runs,  # noqa: E402
                             mcnemar_exact, paired_keys)

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Pt  # noqa: E402

SLIDES = os.path.join(REPO, "paper/slides")
FIGS = os.path.join(SLIDES, "figs")
OUT = os.path.join(SLIDES, "WPMC2026-presentation.pptx")

BG, BAND, INK = "F9F9F9", "DDDDDD", "0A0A0A"
MUTED, BULLET = "6A6A6A", "C4C4C4"
BLUE, BLUE_SOFT, GREY_SOFT = "1E4E8C", "E3ECF7", "EEEEEE"
TITLE_FONT, BODY_FONT, MONO_FONT = "Commissioner", "Golos Text", "Roboto Mono"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------------------
# Numeros, todos calculados dos dados do repositorio
# ---------------------------------------------------------------------------

def compute_numbers() -> dict:
    N = {}
    runs = load_runs(os.path.join(REPO, "logs"))
    keys = paired_keys(runs)
    comp = compliance_table(runs, keys)
    N["canon"] = comp
    N["canon_tasks"] = len(keys)
    N["canon_rule"] = comp["LLM"]["acr"]["n"]
    N["drop"] = {e: 100 * (1 - sum(int(runs[e][k]["success"]) for k in keys) / len(keys)) for e in runs}
    anom = [k for k in keys if runs["LLM"][k]["anomaly"]]

    def disc(a, b, rr, kk):
        b01 = sum(1 for k in kk if rr[a][k]["semantic_compliant"] == "1" and rr[b][k]["semantic_compliant"] == "0")
        b10 = sum(1 for k in kk if rr[a][k]["semantic_compliant"] == "0" and rr[b][k]["semantic_compliant"] == "1")
        return b01, b10, mcnemar_exact(b01, b10)
    N["mc"] = {(a, b): disc(a, b, runs, anom) for a, b in
               (("LLM", "BASELINE"), ("SLM", "BASELINE"), ("LLM", "DRL"), ("SLM", "DRL"),
                ("BASELINE", "DRL"), ("LLM", "SLM"))}
    lat = latency_stats(runs, keys)
    N["lat"] = lat
    N["window_s"] = 5.0 / (1 - math.exp(-5.0 * 4.0 / 60.0))
    N["window_frac"] = 100 * lat["LLM"]["median"] / 1000 / N["window_s"]
    N["energy_frac"] = 100 * 5.0 * N["canon_tasks"] / (50 * 3600)
    N["task_vs_llm"] = 0.02 * 50 * 3600 / 5.0

    ms = load_runs(os.path.join(REPO, "logs/multiseed"), ["BASELINE", "DRL"])
    mk = paired_keys(ms)
    N["ms"] = compliance_table(ms, mk)
    N["ms_tasks"] = len(mk)
    N["ms_seeds"] = len({k.split(":")[0] for k in mk})
    N["ms_p"] = disc("DRL", "BASELINE", ms, [k for k in mk if ms["DRL"][k]["anomaly"]])[2]

    N["t20"] = json.load(open(os.path.join(REPO, "logs/generalization/table20.json")))["splits"]["seen"]
    N["unseen"] = json.load(open(os.path.join(REPO, "logs/generalization/table4_unseen_s0-9.json")))
    gains = []
    for t in (42, 43, 44):
        run = json.load(open(os.path.join(REPO, f"docs/frontier_l12_s{t}_pooled.json")))
        gains.append(next(v["gain_pp"] for n, v in run["agents"].items() if "_w0p05_" in n))
    N["gamma"] = (min(gains), max(gains))

    slm_rows = sum(1 for _ in open(os.path.join(REPO, "logs/mec_metrics_SLM.csv"))) - 1
    for s in N["unseen"]["seeds"]:
        slm_rows += sum(1 for _ in open(os.path.join(
            REPO, f"logs/generalization/heldout_slm_v3/seed{s}/mec_metrics_SLM.csv"))) - 1
    N["slm_decisions"] = slm_rows
    return N


def pct(k, n):
    v = 100 * k / n
    return "100%" if k == n else f"{v:.1f}%"


def pfmt(p):
    """p-valor no formato do artigo: 0.031, 0.0625, 1.000."""
    if p < 0.01:
        return f"{p:.0e}"
    if p >= 0.9995:
        return "1.000"
    four = f"{p:.4f}"
    return four if four.endswith("5") else f"{p:.3f}"


# ---------------------------------------------------------------------------
# Primitivas de desenho
# ---------------------------------------------------------------------------

MARK = re.compile(r"(\*\*.+?\*\*|\^\^.+?\^\^)")


def add_runs(par, text, size, color=INK, font=BODY_FONT, bold=False):
    """**negrito**  e  ^^destaque azul^^."""
    for seg in MARK.split(text):
        if not seg:
            continue
        r = par.add_run()
        b, c = bold, color
        if seg.startswith("**"):
            seg, b = seg[2:-2], True
        elif seg.startswith("^^"):
            seg, b, c = seg[2:-2], True, BLUE
        r.text = seg
        r.font.size, r.font.name, r.font.bold = Pt(size), font, b
        r.font.color.rgb = rgb(c)


def textbox(slide, x, y, w, h, lines, size=16, color=INK, font=BODY_FONT, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=4):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Pt(2))
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(spacing)
        add_runs(p, line, size, color, font, bold)
    return tb


def bullets(slide, x, y, w, h, items, size=16, gap=7):
    """items: str (nivel 0) ou ('sub', str) (nivel 1). Marcador quadrado cinza."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Pt(2))
    for i, it in enumerate(items):
        lvl, text = (1, it[1]) if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i:
            p.space_before = Pt(gap if lvl == 0 else 2)
        pPr = p._p.get_or_add_pPr()
        indent = Pt(18 if lvl == 0 else 16)
        pPr.set("marL", str(int(indent + (Pt(20) if lvl else 0))))
        pPr.set("indent", str(-int(indent)))
        buClr = pPr.makeelement(qn("a:buClr"), {})
        clr = buClr.makeelement(qn("a:srgbClr"), {"val": BULLET})
        buClr.append(clr)
        pPr.append(buClr)
        pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "70000"}))
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))
        add_runs(p, text, size if lvl == 0 else size - 2, INK if lvl == 0 else "333333")
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    # Sem o estilo padrao do python-pptx: ele traz sombra e cores do tema, que a
    # referencia nao tem. Preenchimento, contorno e cor do texto ja sao explicitos.
    style = s._element.find(qn("p:style"))
    if style is not None:
        s._element.remove(style)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def box(slide, x, y, w, h, lines, fill=GREY_SOFT, line=None, size=13, color=INK, bold=False,
        align=PP_ALIGN.CENTER, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE, rounded=True):
    s = rect(slide, x, y, w, h, fill, line, shape=MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Pt(6))
    for m in ("margin_top", "margin_bottom"):
        setattr(tf, m, Pt(3))
    for i, l in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        add_runs(p, l, size, color, font, bold)
    return s


def line(slide, x1, y1, x2, y2, color=INK, w=1.0, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def image(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, Pt(x), Pt(y), Pt(w) if w else None, Pt(h) if h else None)


def table(slide, data, x, y, w, colw, row_h=22, size=11, header=True, fills=None, bold_cols=(),
          aligns=None):
    rows, cols = len(data), len(data[0])
    gt = slide.shapes.add_table(rows, cols, Pt(x), Pt(y), Pt(w), Pt(row_h * rows))
    t = gt.table
    t.first_row = header
    t.horz_banding = False
    for j, cw in enumerate(colw):
        t.columns[j].width = Pt(cw)
    for i in range(rows):
        t.rows[i].height = Pt(row_h)
        for j in range(cols):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Pt(5)
            c.margin_top = c.margin_bottom = Pt(2)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            f = (fills or {}).get((i, j)) or ("D6D6D6" if header and i == 0 else ("FFFFFF" if i % 2 else "F3F3F3"))
            c.fill.solid(); c.fill.fore_color.rgb = rgb(f)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (aligns[j] if aligns else PP_ALIGN.LEFT)
            add_runs(p, str(data[i][j]), size, INK, BODY_FONT, bold=(header and i == 0) or j in bold_cols)
    return gt


# ---------------------------------------------------------------------------
# Moldura dos slides (geometria medida no slide 5 da referencia, 720 x 405 pt)
# ---------------------------------------------------------------------------

def content_slide(prs, num, section, title, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0, 383, 720, 22, BAND)
    line(s, 360, 20, 720, 20, INK, 1.0)
    line(s, 22, 205, 22, 405, INK, 1.0)
    textbox(s, 655, 386, 50, 16, str(num), 10, INK, align=PP_ALIGN.RIGHT)
    if section:
        textbox(s, 30, 9, 300, 14, section.upper(), 8.5, "8A8A8A")
    if title:
        textbox(s, 40, 34, 640, 40, title, 26, INK, align=PP_ALIGN.CENTER)
    s.notes_slide.notes_text_frame.text = " ".join(notes.split())
    return s


def banner(slide, text, y=352, x=60, w=600):
    box(slide, x, y, w, 26, text, fill=BLUE_SOFT, size=14, color=BLUE, bold=True)


def logos(slide, x, y, w, h):
    files = sorted(glob.glob(os.path.join(SLIDES, "logos", "*.png")) +
                   glob.glob(os.path.join(SLIDES, "logos", "*.jpg")))
    if not files:
        box(slide, x, y, w, h, ["UNICAMP / IC logo", "(paper/slides/logos/)"], fill=None, line="BBBBBB",
            size=10, color="8A8A8A")
        return
    slot = w / len(files)
    for i, f in enumerate(files):
        pic = slide.shapes.add_picture(f, Pt(x + i * slot), Pt(y), height=Pt(h))
        if pic.width > Pt(slot - 6):           # cabe na faixa reservada
            ratio = Pt(slot - 6) / pic.width
            pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def build(N) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(720), Pt(405)
    c, u = N["canon"], N["unseen"]
    ue = u["engines"]
    lat = N["lat"]
    g_lo, g_hi = N["gamma"]
    llm_k, slm_k, n_u = ue["LLM"]["k"], ue["SLM"]["k"], u["n"]
    others_max = max(ue[e]["pct"] for e in ("BASELINE", "ORACLE", "DRL", "DRL_ONEHOT"))

    # 1 — capa ---------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 655, 0, 65, 405, BAND)
    line(s, 375, 42, 720, 42, INK, 1.0)
    line(s, 690, 60, 690, 400, INK, 1.0)
    line(s, 0, 362, 380, 362, INK, 1.0)
    textbox(s, 60, 78, 580, 150, ["Semantic Task Orchestration in", "LEO-NTN Multi-Access Edge",
                                  "Computing: A Comparative Study"], 30, INK, TITLE_FONT, bold=True, spacing=0)
    textbox(s, 60, 248, 420, 22, "**Octávio B. Nascimento** and **Carlos A. Astudillo**", 14)
    textbox(s, 60, 270, 420, 32, ["Institute of Computing, University of Campinas (UNICAMP)",
                                  "Campinas, Brazil"], 11, "333333", spacing=1)
    textbox(s, 60, 306, 420, 52, ["**WPMC 2026** — 29th International Symposium on",
                                  "Wireless Personal Multimedia Communications",
                                  "November 8–11, 2026 · Natal, Brazil"], 11, "333333", spacing=1)
    logos(s, 480, 255, 160, 70)
    textbox(s, 655, 386, 50, 16, "1", 10, INK, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = " ".join("""
Good morning. This work compares four ways of orchestrating tasks on a low-Earth-orbit
edge-computing constellation: a greedy heuristic, a deep reinforcement learning agent, a small
language model running on board, and a large language model in the cloud. What makes the
comparison different is the metric. We do not ask which engine is fastest or which one
executes the most tasks; we ask which one respects the rules attached to each task, such as
data-protection law or a hardware-failure protocol. Along the way we will see that the
first, intuitive answer, "the language model wins", needed two corrections suggested by our
reviewers, and that the final answer is more specific and, we think, more useful.""".split())

    # 2 — agenda -------------------------------------------------------------
    s = content_slide(prs, 2, "", "Agenda", """
The talk follows the usual path. First the motivation and the gap. Then the system model:
the constellation, the workload and the semantic constraints. Then the four orchestration
engines and how each one sees a rule. Then the evaluation method, which matters here because
all engines face exactly the same task sequence. The results come in three steps: a single
canonical run, a multi-seed replication, and an information-parity test with rules no engine
was designed for. We close with when to use each paradigm, the limitations and future work.""")
    items = ["Introduction", "System Model", "Orchestration Engines", "Evaluation Method",
             "Results", "Discussion & Conclusion"]
    for i, it in enumerate(items):
        col, row = i % 3, i // 3
        x, y = 70 + col * 205, 120 + row * 110
        box(s, x, y, 52, 70, str(i + 1), fill=BAND, size=26, rounded=False)
        textbox(s, x + 62, y, 135, 70, it, 16, anchor=MSO_ANCHOR.MIDDLE)

    # 3 — motivation ---------------------------------------------------------
    s = content_slide(prs, 3, "1 · Introduction", "Motivation", """
Low-Earth-orbit constellations are becoming compute infrastructure: each satellite can host
multi-access edge computing tasks. Every time a task arrives, the satellite must decide in
real time whether to process it, forward it to a peer over the inter-satellite link, or drop
it, under battery limits driven by the eclipse cycle and limited memory. On top of these
resource constraints, real tasks increasingly carry rules. A GDPR-covered task may only be
processed in Europe. A sovereignty rule pins data to its country. A hardware-failure flag
means the task must not run at all. In the example on the right, a task arrives in Brazil
with a GDPR constraint: processing it locally violates the rule; the compliant decision is to
forward it to the European satellite, and to drop it only if Europe is unavailable.""")
    bullets(s, 45, 95, 330, 270, [
        "LEO constellations as **edge-computing infrastructure** (MEC)",
        "Each satellite decides in real time: **process, forward via ISL, or drop**",
        "Under battery (orbital eclipse) and RAM constraints",
        "Tasks increasingly carry **semantic constraints**:",
        ("sub", "data-protection law (e.g. GDPR)"),
        ("sub", "national data sovereignty"),
        ("sub", "hardware-failure protocols"),
        "A violation is not a performance loss — it is a **compliance failure**",
    ], size=15)
    box(s, 400, 92, 285, 62, ["**Incoming task** · origin BRAZIL · 500 MB · 60 s",
                              "^^Constraint: GDPR — process only in EUROPE^^"], fill="FFFFFF", line="BBBBBB",
        size=12)
    opts = [("Process on SAT-1 (Brazil)", "violates the rule", "F3E3E2", "B0413E"),
            ("Forward via ISL to SAT-2 (Europe)", "compliant", BLUE_SOFT, BLUE),
            ("Drop the task", "only if Europe is unavailable", GREY_SOFT, MUTED)]
    for i, (a, b, f, col) in enumerate(opts):
        y = 182 + i * 58
        box(s, 420, y, 245, 46, [f"**{a}**", b], fill=f, size=12, color=col)
    line(s, 542, 154, 542, 182, "9A9A9A", 1.2, arrow=True)

    # 4 — problem and gap ----------------------------------------------------
    s = content_slide(prs, 4, "1 · Introduction", "Problem and Gap", """
Deep reinforcement learning has been applied to satellite edge scheduling with good results,
but it optimises numerical metrics over numerical feature vectors; a regulatory rule is
invisible unless someone engineers it into the state. Large language models have been shown
to allocate wireless resources and to schedule competitively, but they are slow, which is why
Lee and Park propose hybrid designs. Edge intelligence suggests putting compact language
models on the node itself. The gap is that these threads are studied separately, each
optimising time, energy or throughput. No study puts a heuristic, a DRL agent, an embedded
small model and a cloud model in the same environment with compliance as the primary metric.
So our question is simple: which paradigm respects semantic constraints, and at what cost?""")
    bullets(s, 50, 92, 620, 200, [
        "**DRL schedulers** optimise time, energy or throughput over numeric features — rules are invisible unless engineered into the state",
        "**LLMs** can allocate wireless resources, but pay in latency (Lee & Park, 2024)",
        "**Edge intelligence** suggests compact language models on the node itself",
        "**Gap:** no study compares heuristic, DRL, embedded SLM and cloud LLM in one environment, with **semantic compliance** as the primary metric",
    ], size=15, gap=9)
    box(s, 70, 300, 580, 46, "Which orchestration paradigm respects semantic constraints — and at what cost?",
        fill=BLUE_SOFT, line=BLUE, size=17, color=BLUE, bold=True)

    # 5 — contributions ------------------------------------------------------
    s = content_slide(prs, 5, "1 · Introduction", "Contributions", f"""
Four contributions. First, a comparative framework: four paradigms in the same orbital
scenario, facing the identical task sequence, so differences can be attributed to the
paradigm. Second, an analytical ceiling for the DRL agent: if it only observes that some rule
is present, its expected compliance is at most one third, and twenty seeds confirm it sits
there. Third, and this is the core of the camera-ready version, an information-parity
evaluation: a hand-written rule table and a DRL agent that is told the rule type match the
language models on the known rules, but on rules none of them was designed for only the
language models remain compliant. Fourth, a provenance audit that labels every simulation
parameter as taken from the literature or chosen by us. Every contrast comes with a paired
exact test and a 95% interval.""")
    contrib = [
        "**Comparative framework** — 4 paradigms, same orbital scenario, identical task sequences",
        "**DRL ceiling** — a binary anomaly flag caps expected compliance at 1/3 (analytical, confirmed over 20 seeds)",
        f"**Information parity + unseen rules** — rule table and informed DRL match the LMs on known rules; only LMs generalize (SLM {ue['SLM']['pct']:.0f}%, LLM {ue['LLM']['pct']:.0f}%)",
        "**Parameter provenance audit** — each parameter tagged as literature-grounded or design choice",
    ]
    for i, t in enumerate(contrib):
        y = 88 + i * 62
        box(s, 55, y, 44, 50, str(i + 1), fill=BAND, size=22, rounded=False)
        textbox(s, 110, y, 560, 50, t, 15, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, 55, 340, 610, 22, "Every contrast reported with a paired exact test and a 95% interval",
            12, MUTED, align=PP_ALIGN.CENTER)

    # 6 — architecture -------------------------------------------------------
    s = content_slide(prs, 6, "2 · System Model", "Network Architecture", """
The scenario has three low-Earth-orbit satellites acting as edge servers, each assigned to a
region: SAT-1 to Brazil, SAT-2 to Europe and SAT-15 to the United States. Each has 4,096
megabytes of memory. Tasks reach the constellation through ground stations, and the
inter-satellite links are assumed always available, so a drop is never a link failure: it is
either resource exhaustion or a semantic decision. The envelopes show regular tasks and tasks
carrying one of the three constraint types: GDPR, data sovereignty and hardware fault. The
four engines are plugged into the same simulator in turn; the constellation, the orbit and
the task stream are identical for all of them.""")
    image(s, os.path.join(REPO, "architecture-diagram.png"), 105, 78, w=510)
    textbox(s, 60, 362, 600, 18, "3 regional satellites · 4,096 MB RAM each · inter-satellite links always available",
            11, MUTED, align=PP_ALIGN.CENTER)

    # 7 — workload -----------------------------------------------------------
    s = content_slide(prs, 7, "2 · System Model", "Workload and Semantic Constraints", f"""
Tasks arrive as a Poisson process at four per minute nominal; because the simulator resolves
arrivals on a five-second step, the effective rate is about 3.4 per minute, a mean of
{N['window_s']:.1f} seconds between tasks. Each task holds 500 megabytes for sixty seconds.
Ten percent of tasks carry one of three constraints, chosen uniformly. A constraint arrives as
a short token from a closed vocabulary of three, for example "restricao gdpr europa"; the
language models see that token verbatim, the DRL agent only sees that some constraint exists.
The table shows the compliant decision for each rule. The same predicate scores every engine.
The canonical run lasts 1,200 seconds and contains {N['canon_tasks']} tasks, {N['canon_rule']} of
them with a rule — a number we will come back to.""")
    bullets(s, 45, 92, 300, 270, [
        "Poisson arrivals, λ = 4 tasks/min nominal (≈3.4 effective on the 5 s step)",
        "Each task: 500 MB RAM held for 60 s",
        "**10%** of tasks carry one constraint, drawn uniformly",
        "Constraint = token from a **closed vocabulary of 3**",
        "One predicate scores every engine",
        f"Canonical run: 1,200 s, **{N['canon_tasks']} tasks, {N['canon_rule']} with a rule**",
    ], size=14)
    table(s, [["Constraint", "Compliant decision"],
              ["GDPR (EU privacy)", "route to EUROPE (drop only if unavailable)"],
              ["Data sovereignty (Brazil)", "route to BRAZIL (drop only if unavailable)"],
              ["Critical hardware failure", "drop the task"]],
          360, 120, 320, [130, 190], row_h=40, size=12)

    # 8 — battery & provenance -----------------------------------------------
    s = content_slide(prs, 8, "2 · System Model", "Battery, Eclipse and Parameter Provenance", """
Battery follows an orbital eclipse model: solar input varies sinusoidally with a ninety-minute
period; in eclipse the battery drains, in sunlight it recharges. We set the eclipse to about
53 percent of the orbit, above the roughly 35 minutes reported in the literature, on
purpose: a pessimistic setting that creates real battery pressure inside a twenty-minute
window. Satellites refuse new tasks below 20 percent state of charge. The table is our
provenance audit: each parameter is marked as grounded in the literature or as our design
choice, so you can separate validated premises from modelling decisions. One honest note:
the workload is light, memory occupancy stays at 12 to 14 percent, so memory never binds —
we return to this in the limitations.""")
    bullets(s, 40, 92, 250, 270, [
        "Solar input sinusoidal, orbit **T = 90 min**",
        "Eclipse ⇒ drain; sunlight ⇒ recharge",
        "Eclipse ≈**53%** of the orbit (literature ≈35 min): **pessimistic on purpose**",
        "Satellites refuse tasks below **20% SoC**",
    ], size=14)
    lit, des = BLUE_SOFT, "EDEDED"
    prov = [["Parameter", "Value", "Source"],
            ["Orbital period", "90 min", "Literature"],
            ["Eclipse fraction", "≈53%", "Design (stress)"],
            ["Battery threshold", "20% SoC", "Design"],
            ["SLM latency (NPU)", "50 ms", "Literature"],
            ["Task arrival", "Poisson, 4/min", "Lit. model · design rate"],
            ["RAM total / task", "4,096 / 500 MB", "Design"],
            ["Task hold time", "60 s", "Design"],
            ["Anomaly rate", "10%", "Design"]]
    fills = {(i, 2): (lit if prov[i][2].startswith("Lit") else des) for i in range(1, len(prov))}
    table(s, prov, 305, 88, 380, [130, 100, 150], row_h=30, size=11.5, fills=fills)

    # 9 — four paradigms -----------------------------------------------------
    s = content_slide(prs, 9, "3 · Orchestration Engines", "Four Orchestration Paradigms", """
These are the four engines. The baseline is a greedy heuristic: it filters by region, battery
and memory and picks the satellite with the most charge; it never looks at the rule. The DRL
engine is a Deep Q-Network trained with Stable Baselines 3; its state contains a single bit
saying whether some rule is present, not which one. The small language model is Gemma 4 with
26 billion parameters, standing in for an on-board model on a neural processing unit; we
charge it a 50-millisecond benchmark latency. The large language model is Gemini 3.1
Flash-Lite in the cloud, reached through a ground station. The bottom row is the key
difference: what each engine actually sees of the rule. Everything else — tasks, fleet state
and scoring — is identical.""")
    cards = [("BASELINE", "Greedy heuristic", "region → battery > 20% → RAM; highest battery wins", "On board", "nothing", False),
             ("DRL", "SB3 DQN, MLP 64×64", "trained 50k steps", "On board", "1 bit: a rule is present", False),
             ("SLM", "Gemma 4 26B", "emulated on-device NPU, 50 ms", "On board", "full rule text", True),
             ("LLM", "Gemini 3.1 Flash-Lite", "cloud, via ground station", "Ground / cloud", "full rule text", True)]
    for i, (name, m1, m2, where, sees, lm) in enumerate(cards):
        x = 40 + i * 162
        box(s, x, 85, 152, 34, f"**{name}**", fill=BLUE if lm else "7A7A7A", size=16, color="FFFFFF", rounded=False)
        box(s, x, 119, 152, 130, [f"**{m1}**", m2, "", f"Runs: {where}"], fill="FFFFFF", line="CCCCCC",
            size=12, rounded=False, anchor=MSO_ANCHOR.TOP)
        box(s, x, 255, 152, 50, ["Sees of the rule:", f"**{sees}**"], fill=BLUE_SOFT if lm else GREY_SOFT,
            size=12, color=BLUE if lm else INK, rounded=False)
    banner(s, "Same tasks, same fleet state, same scoring predicate for every engine", y=322)

    # 10 — DRL ---------------------------------------------------------------
    s = content_slide(prs, 10, "3 · Orchestration Engines", "DRL Formulation and Its Ceiling", """
The DRL agent sees a thirteen-dimensional state: the task region, one bit saying whether a
rule is present, and battery, free memory and solar status for each satellite. It chooses one
of the three satellites or drops the task. Each arrival is one decision, so the discount
factor is zero, a contextual bandit, appropriate here because orbital dynamics are slow
compared with a single decision. During training only, the hidden rule type shapes the
reward, but it is never part of the observation. That gives a clean ceiling: when the bit is
one, the agent cannot know whether the right action is Europe, Brazil or drop. A blind policy
commits to one action, and with three equally likely rules it is right for at most one of
them. So expected compliance is at most one third, regardless of training budget.""")
    box(s, 40, 88, 220, 70, ["**State s ∈ [0,1]¹³**", "region one-hot (3) · rule flag h (1)",
                             "per satellite: battery, RAM, solar (9)"], fill="FFFFFF", line="AAAAAA", size=11.5)
    box(s, 290, 88, 140, 70, ["**DQN**", "MLP 64×64", "γ = 0 · 50k steps"], fill=GREY_SOFT, size=12)
    box(s, 460, 88, 220, 70, ["**Action**", "SAT-1 · SAT-2 · SAT-15 · DROP"], fill="FFFFFF", line="AAAAAA", size=12)
    line(s, 260, 123, 290, 123, INK, 1.3, arrow=True)
    line(s, 430, 123, 460, 123, INK, 1.3, arrow=True)
    textbox(s, 40, 170, 300, 22, "**Reward**", 14)
    bullets(s, 40, 192, 300, 160, [
        "+1 + 0.1·battery: valid route",
        "+1: drop on a hardware fault",
        "−0.5 forced drop · −1 voluntary drop",
        "−2 invalid action",
        "Rule type shapes the reward **in training only** — never observed",
    ], size=13, gap=4)
    box(s, 360, 175, 320, 170, ["**Why expected ACR ≤ 1/3**", "",
                                "h = 1 hides which rule applies:",
                                "GDPR → EUROPE · sovereignty → BRAZIL · fault → DROP", "",
                                "A blind policy commits to one action; with three equiprobable rules it is right for at most one."],
        fill=BLUE_SOFT, line=BLUE, size=13, color=INK)

    # 11 — prompt ------------------------------------------------------------
    s = content_slide(prs, 11, "3 · Orchestration Engines", "Language Models: One Shared Prompt", """
Both language models receive exactly the same prompt, shown here abridged: a role, the task
with its region, memory and constraint token, the live fleet state, and four ordered routing
rules. For a rule outside that list — the unseen rules we test later — one sentence stating
the constraint is appended to the task line, and that sentence is the only information the
model gets. Both run at temperature zero in JSON mode and answer with the identifier of the
target satellite, or null to drop, read by the same parser. One difference is forced by the
models themselves: Gemma reasons before answering, so its output budget is 8,192 tokens and
only its final answer is read. A truncated or malformed answer counts as a failure, never as
a drop.""")
    prompt = ["You are a centralized satellite network",
              "orchestrator managing a LEO constellation.",
              "Output ONLY valid JSON.", "",
              "TASK: id={id}, region={region},",
              "      ram={ram} MB, anomaly=\"{anomaly}\"", "",
              "AVAILABLE SATELLITES:",
              "  SAT 1 | region=BRAZIL | battery_pct=64%",
              "  ... (one line per satellite) ...", "",
              "ROUTING RULES (apply in order):",
              " 1. GDPR anomaly → EUROPE satellite",
              " 2. Sovereignty anomaly → required country",
              " 3. Hardware failure → drop (null)",
              " 4. No anomaly → task region, battery > 20%", "",
              "Output: {\"satellite_id\": <id or null>,",
              "         \"reason\": \"one short sentence\"}"]
    box(s, 38, 84, 330, 288, prompt, fill="FFFFFF", line="BBBBBB", size=10.5, font=MONO_FONT,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, rounded=False)
    bullets(s, 385, 90, 300, 280, [
        "**One prompt**, shared verbatim by SLM and LLM",
        "Rule type arrives **as text**; an unseen rule adds one sentence stating it",
        "temperature = 0, JSON mode",
        "Answer = **satellite_id** (or null = drop), same parser",
        "Gemma reasons first: 8,192-token budget; only the final answer is read",
        "Truncated or malformed answer = **failure**, not a drop",
    ], size=14)

    # 12 — protocol ----------------------------------------------------------
    s = content_slide(prs, 12, "4 · Evaluation Method", "Metrics and Statistical Protocol", f"""
The primary metric is anomaly compliance, ACR: among tasks that carry a rule, the fraction
handled as the rule requires, including mandatory drops. Semantic compliance, SCR, is the same
over all tasks. We also report execution and drop rates and decision latency; energy per
decision is a declared model, not a measurement. The design is paired: for a given seed every
engine sees exactly the same arrivals, regions and rules, so we compare engines task by task
with McNemar's exact test and report Wilson 95% intervals. The table shows the three
experiments: one canonical run, twenty seeds for the deterministic engines on the known rules,
and ten seeds with all six engines on the unseen rules.""")
    bullets(s, 38, 88, 320, 150, [
        "**ACR**: compliance on rule-carrying tasks — primary KPI",
        "**SCR**: compliance over all tasks",
        "Execution / drop rate · decision latency",
        "Energy per decision: **declared model**, not measured",
    ], size=13.5, gap=5)
    ms_n = N["ms"]["BASELINE"]["acr"]["n"]
    table(s, [["Experiment", "Seeds", "Rule tasks", "Engines"],
              ["Canonical run", "1", str(N["canon_rule"]), "4"],
              ["Known rules", str(N["ms_seeds"]), str(ms_n), "4 deterministic"],
              ["Unseen rules", str(len(u["seeds"])), str(n_u), "all 6 (paired)"]],
          38, 250, 320, [95, 45, 80, 100], row_h=24, size=11.5,
          aligns=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.LEFT])
    box(s, 395, 86, 285, 50, ["**Seed s** → identical task sequence", "arrivals · regions · rules"],
        fill="FFFFFF", line="AAAAAA", size=12.5)
    line(s, 537, 136, 537, 158, INK, 1.2, arrow=True)
    for i, e in enumerate(["BASELINE", "ORACLE", "DRL", "DRL-OH", "SLM", "LLM"]):
        col, row = i % 3, i // 3
        box(s, 395 + col * 97, 160 + row * 34, 91, 28, e, fill=BLUE_SOFT if e in ("SLM", "LLM") else GREY_SOFT,
            size=11.5, color=BLUE if e in ("SLM", "LLM") else INK, bold=True)
    line(s, 537, 226, 537, 248, INK, 1.2, arrow=True)
    box(s, 395, 250, 285, 50, ["Paired outcome per task", "**McNemar exact · Wilson 95% CI**"],
        fill=BLUE_SOFT, line=BLUE, size=12.5, color=BLUE)

    # 13 — canonical run -----------------------------------------------------
    mc = N["mc"]
    s = content_slide(prs, 13, "5 · Results", "First Look: The Canonical Run", f"""
Here is the canonical run. Both language models handle all {N['canon_rule']} rule-carrying
tasks correctly, the heuristic {c['BASELINE']['acr']['k']}, the DRL agent {c['DRL']['acr']['k']}.
It is tempting to stop here and say the language models win. But look at the intervals: with
eight tasks one decision moves the rate by 12.5 points, and the eight-of-eight carries a 95%
interval from about 68 to 100 percent. The paired test confirms the caution: only language
model versus baseline is significant, p equals {pfmt(mc[('LLM', 'BASELINE')][2])}; language
model versus DRL is not, p equals {pfmt(mc[('LLM', 'DRL')][2])}. On drops, the language models
drop only three tasks, all mandatory hardware-failure discards. So the picture is promising,
but two questions remain: is it robust, and is the comparison fair?""")
    image(s, os.path.join(FIGS, "fig_acr_canonical.png"), 32, 82, w=385)
    bullets(s, 430, 88, 260, 260, [
        f"Both language models: **{c['LLM']['acr']['k']}/{c['LLM']['acr']['n']}**",
        f"Heuristic {c['BASELINE']['acr']['k']}/{N['canon_rule']}, DRL {c['DRL']['acr']['k']}/{N['canon_rule']}",
        "But **n = 8**: one task moves ACR by 12.5 pp — intervals overlap",
        f"Significant only LM vs BASELINE ({mc[('LLM', 'BASELINE')][0]}/{mc[('LLM', 'BASELINE')][1]}, p = {pfmt(mc[('LLM', 'BASELINE')][2])}); LM vs DRL p = {pfmt(mc[('LLM', 'DRL')][2])}",
        f"Drops: BASELINE {N['drop']['BASELINE']:.0f}%, DRL {N['drop']['DRL']:.1f}%, SLM/LLM {N['drop']['LLM']:.1f}% (mandatory discards only)",
    ], size=13, gap=6)
    banner(s, "Promising — but is it robust, and is the comparison fair?", y=318)

    # 14 — multiseed ---------------------------------------------------------
    b, d = N["ms"]["BASELINE"]["acr"], N["ms"]["DRL"]["acr"]
    s = content_slide(prs, 14, "5 · Results", "Does the DRL Gain Replicate?", f"""
First question: robustness. For the deterministic engines we repeated the experiment over
{N['ms_seeds']} seeds, {N['ms_tasks']:,} paired tasks and {b['n']} with a rule. The DRL edge
over the heuristic disappears: {pct(d['k'], d['n'])} against {pct(b['k'], b['n'])}, p equals
{N['ms_p']:.3f}. The 37.5 percent of the single run was chance. More interesting, the DRL agent
sits right on the one-third bound we derived analytically. That is the point of this slide:
the DRL agent is not badly trained; it is limited by what it observes. It cannot tell which
rule applies, so no amount of training will push it past one third. That leads directly to
the second question: what happens if we give it the rule?""")
    image(s, os.path.join(FIGS, "fig_multiseed.png"), 40, 82, w=345)
    bullets(s, 400, 92, 285, 240, [
        f"{N['ms_seeds']} seeds · {N['ms_tasks']:,} paired tasks · {b['n']} with a rule",
        f"DRL **{pct(d['k'], d['n'])}** vs BASELINE **{pct(b['k'], b['n'])}** — p = {N['ms_p']:.3f}",
        "The single-run DRL edge (37.5%) was chance",
        "DRL sits on the **1/3 bound** predicted analytically",
        "⇒ The limit is **what DRL observes**, not DRL itself",
    ], size=14)
    banner(s, "Next: give DRL — and a rule table — the rule itself", y=318)

    # 15 — information parity ------------------------------------------------
    s = content_slide(prs, 15, "5 · Results", "Information Parity and Unseen Rules", """
This is the question our third reviewer asked: is the comparison fair if the DRL agent sees
less? We added two control arms. ORACLE is a hand-coded table of the three known rules with
the heuristic as fallback. DRL-OH is the same DQN retrained with a one-hot encoding of the
rule type, so it knows exactly which rule applies. If language models only win because they
see more, these controls should match them. Then we added a second rule set that none of the
engines was designed for: US HIPAA and ITAR, Brazilian LGPD, a paraphrased EU data-residency
rule, and two new hardware faults, delivered only as a sentence. One third of them are
discard rules, as before, so the blind-policy bound stays the same and the two rule sets are
directly comparable.""")
    table(s, [["Engine", "What it knows about the rule", "Role"],
              ["BASELINE", "nothing", "reference"],
              ["ORACLE", "hand-coded table of the 3 known rules", "control: enumerated rules"],
              ["DRL", "1 bit: a rule is present", "original arm"],
              ["DRL-OH", "one-hot of the 3 known rules", "control: informed DRL"],
              ["SLM", "rule as natural-language text", "language model"],
              ["LLM", "rule as natural-language text", "language model"]],
          32, 86, 405, [70, 195, 140], row_h=34, size=11.5, bold_cols=(0,),
          fills={(i, 0): BLUE_SOFT for i in (5, 6)})
    box(s, 450, 86, 238, 238, ["**Unseen rule set**", "none of the engines was designed for it", "",
                               "US HIPAA → USA", "US ITAR → USA", "Brazilian LGPD → BRAZIL",
                               "EU data residency (paraphrase) → EUROPE", "2 new hardware faults → DROP", "",
                               "delivered only as **one sentence**", "1/3 discard rules: **same blind bound**"],
        fill="FFFFFF", line=BLUE, size=12, rounded=False)

    # 16 — result ------------------------------------------------------------
    s = content_slide(prs, 16, "5 · Results", "Design-Time vs Unseen Rules", f"""
This is the key result. Light bars are the known rules, dark bars the unseen ones. On the
known rules, both controls reach 100 percent: at this scale, enumerating the rules suffices,
and you do not need a language model. On the unseen rules the picture inverts. The rule table
falls to {ue['ORACLE']['pct']:.1f} percent; in fact it takes exactly the same decision as the
heuristic on every task. It does not fail loudly: it silently stops noticing the obligation.
The informed DRL agent falls to {ue['DRL_ONEHOT']['pct']:.1f}; with its rule vector zeroed it
has nothing to act on. Only the language models stay high: the LLM handles {llm_k} of {n_u}
paired tasks, the SLM {slm_k} of {n_u}, both with p below ten to the minus eight against the
heuristic.""")
    image(s, os.path.join(FIGS, "fig_generalization.png"), 55, 76, w=610)
    textbox(s, 50, 343, 620, 38, [
        f"Known rules: **enumerating them suffices**.  Unseen rules: **only the language models stay compliant** — "
        f"LLM {llm_k}/{n_u}, SLM {slm_k}/{n_u} (p < 10⁻⁸ vs BASELINE); tables and DRL ≤ {others_max:.0f}%"],
        13, align=PP_ALIGN.CENTER)

    # 17 — inside ------------------------------------------------------------
    nr, nf = u["n_region"], u["n_fault"]
    lvs = u["llm_vs_slm"]
    s = content_slide(prs, 17, "5 · Results", "Inside the Unseen-Rule Result", f"""
Splitting the unseen tasks by obligation explains the numbers. On the {nr} tasks that require
routing to a specific region, the heuristic gets {ue['BASELINE']['region']} right by geographic
coincidence, the SLM {ue['SLM']['region']}, the LLM all {ue['LLM']['region']}. The DRL agent
gets every fault right and almost no region rule, because it simply drops every flagged task.
What about the SLM's five misses? We replayed the exact prompts: in all five, Gemma answered
with a JSON list wrapping the object, and each one names a satellite in the required region.
Under the protocol we fixed before running, that is a format failure, and we count it. The
consequence is that LLM and SLM are not separable here, five to zero discordant pairs, p
equals {lvs['p']:.4f}. The distinction is language model versus enumerated rules, not which
language model.""")
    rows = [["Engine", f"Region {nr}", f"Fault {nf}", f"Total {n_u}"]]
    for e, lab in (("BASELINE", "BASELINE"), ("ORACLE", "ORACLE"), ("DRL", "DRL"), ("DRL_ONEHOT", "DRL-OH"),
                   ("SLM", "SLM"), ("LLM", "LLM")):
        rows.append([lab, str(ue[e]["region"]), str(ue[e]["fault"]), str(ue[e]["k"])])
    table(s, rows, 32, 86, 300, [76, 80, 72, 72], row_h=30, size=12, bold_cols=(0,),
          fills={(i, j): BLUE_SOFT for i in (5, 6) for j in range(4)},
          aligns=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER])
    bullets(s, 350, 86, 335, 170, [
        "DRL drops every flagged task: right on faults, wrong on regions",
        "Rule table = heuristic on every task (0 discordant pairs)",
        "All 5 SLM misses: **a JSON list instead of an object** — each naming the right region — counted as failures",
        f"LLM vs SLM: {lvs['b01']}/{lvs['b10']} discordant, p = {pfmt(lvs['p'])} → **not separable**",
    ], size=12.5, gap=5)
    box(s, 350, 262, 335, 34, '[{"satellite_id": 1, "reason": "...BRAZIL region."}]', fill="FFFFFF",
        line="BBBBBB", size=10, font=MONO_FONT, rounded=False)
    banner(s, "A language model vs enumerated rules — not which language model", y=318)

    # 18 — cost --------------------------------------------------------------
    s = content_slide(prs, 18, "5 · Results", "The Cost: Latency and Energy", f"""
What does a language model cost? Latency first. The heuristic and the DRL agent decide in
microseconds; the SLM is charged its 50-millisecond NPU benchmark; the cloud LLM has a median
of {lat['LLM']['median']:.0f} milliseconds, with an interquartile range from
{lat['LLM']['p25']:.0f} to {lat['LLM']['p75']:,.0f}. The mean is higher, but that comes from the
free-tier rate-limit back-off, not inference. The relevant budget here is the time between
tasks, {N['window_s']:.1f} seconds, and the LLM's median uses {N['window_frac']:.1f} percent of
it. Energy, under our declared model, differs by orders of magnitude between engines, yet all
the LLM decisions of a run add up to {N['energy_frac']:.2f} percent of the smallest battery,
and executing one task costs as much as about {N['task_vs_llm']:.0f} LLM decisions. The
binding cost is latency and dependence on a ground link, not energy.""")
    image(s, os.path.join(FIGS, "fig_latency.png"), 95, 76, w=530)
    box(s, 60, 292, 290, 54, ["**Latency**", f"LLM median {lat['LLM']['median']:.0f} ms (IQR {lat['LLM']['p25']:.0f}–{lat['LLM']['p75']:,.0f}) = {N['window_frac']:.1f}% of the {N['window_s']:.1f} s between tasks"],
        fill="FFFFFF", line="CCCCCC", size=11.5)
    box(s, 370, 292, 290, 54, ["**Energy** (declared model)", f"all LLM decisions of a run = {N['energy_frac']:.2f}% of the smallest battery; one task ≈ {N['task_vs_llm']:.0f} LLM decisions"],
        fill="FFFFFF", line="CCCCCC", size=11.5)
    banner(s, "The binding cost is latency and a ground link — not energy", y=352)

    # 19 — when to use -------------------------------------------------------
    s = content_slide(prs, 19, "6 · Discussion & Conclusion", "When to Use Each Paradigm", f"""
Putting the results together gives a practical guide. If every rule is known at design time,
a rule table or a DRL agent that is told the rule type is enough: both reach 100 percent,
decide in microseconds and need no external service. If constraints can arrive that nobody
enumerated — new regulations, paraphrased rules — a language model is the only paradigm that
kept up in our tests. If the decision window is seconds, as here, the cloud LLM is viable; if
the budget is milliseconds, the cloud round trip rules it out and an embedded small model, or
plain rules, is the candidate, provided on-board inference achieves NPU-class latency.""")
    table(s, [["If…", "Use", "Why"],
              ["Every rule is known at design time", "Rule table or informed DRL", "100% on known rules · µs latency · no API"],
              ["New or open-ended constraints arrive", "A language model", f"Only paradigm that generalized (LLM {ue['LLM']['pct']:.0f}%, SLM {ue['SLM']['pct']:.0f}%)"],
              ["Decision window of seconds (as here)", "Cloud LLM is viable", f"≈1 s median vs {N['window_s']:.1f} s between tasks"],
              ["Millisecond budget", "Embedded SLM (or rules)", "A cloud round-trip ≈1 s rules the LLM out"]],
          40, 95, 640, [210, 170, 260], row_h=50, size=13, bold_cols=(1,))

    # 20 — limitations -------------------------------------------------------
    s = content_slide(prs, 20, "6 · Discussion & Conclusion", "Limitations and Preliminary Results", f"""
Five limitations bound these conclusions. The workload is light, so memory never binds; that is
precisely why a single-step agent with gamma zero is appropriate here. We ran a preliminary
sequential study under heavy battery load: an agent with gamma 0.99 and throughput weight 0.05,
the only weight whose gain replicated in all three training seeds, reaches {g_lo:.1f} to
{g_hi:.1f} points higher compliance than the best mix of 24 myopic agents at equal throughput,
and gains nothing under moderate load; the plot shows those three agents above the myopic
frontier. The unseen set has six rules, and they are public regulations, so the models may
have seen them in pre-training; invented rules would separate reading from prior knowledge.
Energy is a declared model, and the SLM is emulated through a remote API with an NPU benchmark
latency.""")
    bullets(s, 35, 86, 345, 290, [
        "**Light load**: RAM 12–14%, never binds → single-step γ = 0 fits this regime",
        f"**Preliminary**, heavy battery load: γ = 0.99 (w = 0.05) gains **+{g_lo:.1f}–{g_hi:.1f} pp** ACR over the best mix of 24 myopic agents — 3/3 training seeds",
        "**Six unseen rules**; public regulations may be in the LMs' pre-training",
        "Energy per decision: **declared model**, not measured",
        "SLM emulated via a remote API with an NPU **benchmark** latency",
    ], size=13, gap=7)
    image(s, os.path.join(FIGS, "fig_frontier.png"), 392, 86, w=298)

    # 21 — conclusion --------------------------------------------------------
    s = content_slide(prs, 21, "6 · Discussion & Conclusion", "Conclusion and Future Work", f"""
To conclude. When every rule is known and enumerated, a rule table or an informed DRL agent
matches the language models, at a fraction of the cost. What separates the language models is
generalisation: on rules no one programmed, the LLM handled {llm_k} of {n_u} tasks and the SLM
{slm_k} of {n_u}, while tables and DRL fell to heuristic level. And the price of a language
model is latency and dependence on a link, not energy. For future work, the natural next step
is a hybrid macro-micro design, with an LLM reconfiguring policy and an SLM or DRL agent
deciding per task; longer windows spanning several eclipses and heavier loads; invented rules
to separate rule reading from prior knowledge; and sequential agents under real resource
pressure.""")
    textbox(s, 45, 82, 300, 24, "**Conclusions**", 16)
    bullets(s, 45, 108, 305, 230, [
        "Known, enumerated rules → a **table or informed DRL** matches the LMs",
        f"Unseen rules → **only LMs stay compliant** (LLM {llm_k}/{n_u}, SLM {slm_k}/{n_u}); tables and DRL ≈ heuristic",
        "The cost of an LM is **latency and link dependence**, not energy",
    ], size=13.5, gap=8)
    textbox(s, 375, 82, 300, 24, "**Future work**", 16)
    bullets(s, 375, 108, 305, 230, [
        "Hybrid macro-micro: LLM reconfigures policy, SLM/DRL decides per task",
        "≥ 2 h windows across several eclipses; heavier loads",
        "Invented rules: rule reading vs prior knowledge",
        "Sequential (γ > 0) agents under resource pressure",
    ], size=13.5, gap=8)
    banner(s, "A language model earns its place when constraints arrive that no one enumerated", y=340)

    # 22 — thanks ------------------------------------------------------------
    s = content_slide(prs, 22, "", "", """
Thank you. The simulator, the data and the scripts that regenerate every number in the paper
are in the public repository on the slide. We are happy to take questions.""")
    textbox(s, 60, 60, 600, 70, "Thanks!", 48, INK, TITLE_FONT, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, 60, 150, 600, 80, ["Questions?", "",
                                  "**Octávio B. Nascimento**: o241327@dac.unicamp.br",
                                  "**Carlos A. Astudillo**: castudillo@unicamp.br"], 14, align=PP_ALIGN.CENTER, spacing=2)
    textbox(s, 60, 245, 600, 20, "Code and data: github.com/wocn-unicamp/CosmicBeats-Simulator", 12, MUTED,
            align=PP_ALIGN.CENTER)
    logos(s, 280, 285, 160, 70)

    # ---------------------------------------------------------------- reserva
    # B1 — parametros
    s = content_slide(prs, "B1", "Backup", "Simulation Parameters and Provenance", """
Backup. The full parameter table of the paper, with the source or rationale of each value.
The battery capacities are nominal: state-of-charge dynamics are in percentage points and do
not scale with capacity, so satellites differ in initial charge (90, 75 and 85 percent) and
orbital phase.""")
    table(s, [["Parameter", "Value", "Source / rationale"],
              ["Orbital period", "90 min", "Typical LEO range (Guo 2024; 3GPP 2022)"],
              ["Eclipse fraction", "≈53%", "Above typical ≈35 min (Sumanth 2019); pessimistic stress"],
              ["Battery threshold", "20% SoC", "Design choice"],
              ["Battery capacities", "75 / 50 / 100 Wh", "Design choice (nominal; initial SoC 90/75/85%)"],
              ["SLM latency (NPU)", "50 ms", "Order of magnitude from on-device benchmarks (Wang 2025)"],
              ["Task arrival", "Poisson, 4/min (3.4 eff.)", "Standard NTN/IoT model (Guo 2024); rate is a design choice"],
              ["RAM total / task", "4,096 / 500 MB", "Design choice"],
              ["Task hold time", "60 s", "Design choice"],
              ["Anomaly rate", "10%", f"Design choice (≈{N['canon_rule']} rule tasks in {N['canon_tasks']})"]],
          40, 84, 640, [125, 155, 360], row_h=29, size=11, bold_cols=(0,))

    # B2 — DQN
    s = content_slide(prs, "B2", "Backup", "DQN Hyperparameters and Reward", """
Backup. The DQN hyperparameters and the reward. Training states are drawn independently:
region uniform, rule present with probability 0.1, battery uniform between 15 and 100 percent,
free memory uniform. With gamma equal to zero the Bellman target reduces to the immediate
reward, so learning is regression on the reward function. One practical pitfall reported in
the paper: loading a model in Stable Baselines 3 silently resets the global Python random seed,
which would have broken the shared task sequence; we restore the seed right after loading.""")
    table(s, [["Hyperparameter", "Value"], ["Policy network", "MLP [64, 64]"], ["Learning rate", "10⁻³"],
              ["Replay buffer", "20,000"], ["Batch size", "128"], ["Discount γ", "0 (single-step)"],
              ["Target update", "hard, every 500 steps"], ["Train frequency", "every 4 steps"],
              ["Exploration", "30% of training, final ε = 0.05"], ["Training steps", "50,000"],
              ["Evaluation", "500 det. episodes / 5,000 steps"], ["Seed", "42"]],
          40, 84, 340, [130, 210], row_h=24, size=11.5)
    box(s, 400, 84, 280, 200, ["**Reward R(s, a)**", "",
                               "+1 + 0.1·b   valid route to eligible satellite",
                               "+1   drop on a hardware-failure task",
                               "−0.5   forced drop (no valid candidate)",
                               "−1   voluntary drop while a candidate exists",
                               "−2   invalid action (region / SoC / RAM)", "",
                               "Hidden rule type selects the compliant target **during training only**"],
        fill="FFFFFF", line="BBBBBB", size=11.5, align=PP_ALIGN.LEFT, rounded=False)

    # B3 — protocolo do SLM
    s = content_slide(prs, "B3", "Backup", "SLM Protocol: v1 → v3 (Camera-Ready)", f"""
Backup, in case someone asks why the SLM numbers differ from the submitted version. In the
submission the SLM had its own prompt, answered with an action and a region translated by code,
and had a 1,024-token output budget; with that protocol it scored 50 percent on the unseen
rules we added during the revision. Replaying its failures showed that in all thirteen
format failures the reasoning had exhausted the budget before any answer was written. For the
camera-ready we gave the SLM the LLM's prompt and output format, a budget of 8,192 tokens, and
a strict reader, and we committed the protocol and the rule for reporting three or ten seeds
before running anything. None of the {N['slm_decisions']} decisions reported was truncated.
If asked about latency: the real API call takes several seconds because Gemma reasons first;
the simulation charges the 50-millisecond NPU benchmark, since the question is the decision,
not the API.""")
    table(s, [["", "Protocol v1 (submission)", "Protocol v3 (camera-ready)"],
              ["Prompt", "SLM-specific template", "identical to the LLM's"],
              ["Answer", "action + region, translated by code", "satellite_id, read as given"],
              ["Output budget", "1,024 tokens", "8,192 tokens"],
              ["Reading", "could recover decisions from reasoning text", "final answer only; malformed = failure"],
              ["Unseen rules (ACR)", "50.0% (3 seeds, measured in revision)", f"{ue['SLM']['pct']:.1f}% ({len(u['seeds'])} seeds)"]],
          40, 84, 640, [140, 250, 250], row_h=34, size=12, bold_cols=(0,))
    bullets(s, 50, 290, 620, 80, [
        "At 1,024 tokens, **13/13** format failures were reasoning cut off before any answer",
        f"Protocol and the 3-vs-10-seed rule **committed before any result**; none of the {N['slm_decisions']} decisions truncated",
    ], size=13, gap=4)

    # B4 — estatistica
    s = content_slide(prs, "B4", "Backup", "Canonical Run: Full KPIs and Paired Tests", """
Backup with the full Table III of the paper and the paired tests of the canonical run. SLM and
LLM take identical decisions on every one of the 69 tasks, so their comparison has zero
discordant pairs. Energy values are declared model constants, not measurements; the latency of
the SLM is the NPU benchmark, and that of the LLM is measured wall-clock.""")
    eng = ["BASELINE", "DRL", "SLM", "LLM"]

    def ci(e, m):
        a = c[e][m]; lo, hi = a["wilson"]
        return f"{pct(a['k'], a['n'])} [{100 * lo:.1f}–{'100' if hi > 0.9999 else f'{100 * hi:.1f}'}]"
    lat_s = {e: (f"{lat[e]['median']:.3f} ms" if lat[e]["median"] < 1 else f"{lat[e]['median']:.0f} ms") for e in eng}
    table(s, [["KPI"] + eng,
              ["SCR"] + [ci(e, "scr") for e in eng],
              ["ACR"] + [ci(e, "acr") for e in eng],
              ["Drop rate"] + [f"{N['drop'][e]:.1f}%" for e in eng],
              ["Median latency"] + [lat_s[e] for e in eng],
              ["Energy / decision*"] + ["0.001 J", "0.005 J", "0.100 J", "5.0 J"]],
          30, 84, 660, [124, 134, 134, 134, 134], row_h=28, size=11, bold_cols=(0,))
    pairs = [("LLM", "BASELINE"), ("SLM", "BASELINE"), ("LLM", "DRL"), ("SLM", "DRL"), ("BASELINE", "DRL"), ("LLM", "SLM")]
    table(s, [["Paired test (rule tasks)", "Discordant", "McNemar p"]] +
          [[f"{a} vs {b_}", f"{mc[(a, b_)][0]}/{mc[(a, b_)][1]}", pfmt(mc[(a, b_)][2])] for a, b_ in pairs],
          30, 262, 400, [190, 100, 110], row_h=17, size=10)
    textbox(s, 450, 262, 240, 40, "*Declared model constants, not measurements", 10, MUTED)
    return prs


def main() -> int:
    N = compute_numbers()
    prs = build(N)
    prs.save(OUT)
    n_main = 22
    total = len(prs.slides)
    assert total == n_main + 4, f"esperava 26 slides (22 + 4 reserva), gerou {total}"
    print(f"[ok] {OUT}: {total} slides ({n_main} principais + {total - n_main} reserva)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
